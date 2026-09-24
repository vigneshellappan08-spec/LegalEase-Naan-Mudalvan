from io import BytesIO
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from fpdf import FPDF

from services.sanitization import (
    sanitize_text,
    split_terms
)


BASE_DIR = Path(__file__).resolve().parent.parent

ASSETS_DIR = BASE_DIR / "assets"


def format_txt(text: str) -> bytes:

    clean_text = sanitize_text(text)

    return clean_text.encode("utf-8")


def format_docx(
    text: str,
    doc_type: str,
    terms: Optional[str] = None,
    logo_bytes: Optional[bytes] = None,
    logo_name: str = "logo.png"
) -> bytes:

    text = sanitize_text(text)

    document = Document()


    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)


    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)


    # Logo

    if logo_bytes:

        temporary_logo = (
            ASSETS_DIR / f"_temporary_{logo_name}"
        )

        temporary_logo.write_bytes(
            logo_bytes
        )

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run()

            run.add_picture(
                str(temporary_logo),
                width=Inches(1.2)
            )

        finally:

            temporary_logo.unlink(
                missing_ok=True
            )


    # Document title

    title_paragraph = document.add_paragraph()

    title_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    title_run = title_paragraph.add_run(
        doc_type.upper()
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)


    # Document content

    for raw_line in text.splitlines():

        line = raw_line.strip()


        if not line:

            document.add_paragraph()

            continue


        paragraph = document.add_paragraph()


        if line.startswith("#"):

            heading = line.lstrip("#").strip()

            run = paragraph.add_run(
                heading
            )

            run.bold = True
            run.font.size = Pt(13)


        elif (
            len(line) >= 3
            and line[0].isdigit()
            and ". " in line[:6]
        ):

            run = paragraph.add_run(line)

            run.bold = True


        elif line.startswith("-"):

            paragraph.style = (
                document.styles["List Bullet"]
            )

            paragraph.add_run(
                line[1:].strip()
            )


        elif line.startswith("*"):

            paragraph.style = (
                document.styles["List Bullet"]
            )

            paragraph.add_run(
                line[1:].strip()
            )


        else:

            paragraph.add_run(line)


    # Terms table

    if terms:

        term_list = split_terms(terms)


        if term_list:

            document.add_paragraph()


            heading = document.add_paragraph()

            heading_run = heading.add_run(
                "Key Terms"
            )

            heading_run.bold = True


            table = document.add_table(
                rows=1,
                cols=2
            )

            table.style = "Table Grid"

            table.alignment = (
                WD_TABLE_ALIGNMENT.CENTER
            )


            header = table.rows[0].cells

            header[0].text = "No."
            header[1].text = "Term"


            for index, term in enumerate(
                term_list,
                start=1
            ):

                cells = table.add_row().cells

                cells[0].text = str(index)
                cells[1].text = term


    # Footer

    footer = (
        section.footer.paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-assisted draft. "
        "Review before use."
    )

    footer_run.italic = True


    output = BytesIO()

    document.save(output)

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(
        self,
        doc_type: str
    ):

        super().__init__()

        self.doc_type = doc_type


    def header(self):

        self.set_font(
            "Times",
            "B",
            14
        )

        self.cell(
            0,
            8,
            self.doc_type.upper(),
            align="C"
        )

        self.ln(12)


    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Times",
            "I",
            8
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-assisted draft. "
            "Review before use.",
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str
) -> bytes:

    clean_text = sanitize_text(text)


    pdf = LegalEasePDF(
        doc_type=doc_type
    )


    pdf.set_auto_page_break(
        auto=True,
        margin=18
    )


    pdf.add_page()


    pdf.set_font(
        "Times",
        size=11
    )


    for raw_line in clean_text.splitlines():

        line = raw_line.strip()


        if not line:

            pdf.ln(5)

            continue


        if line.startswith("#"):

            pdf.set_font(
                "Times",
                "B",
                13
            )

            pdf.multi_cell(
                0,
                7,
                line.lstrip("#").strip()
            )

            pdf.set_font(
                "Times",
                size=11
            )


        elif (
            len(line) >= 3
            and line[0].isdigit()
            and ". " in line[:6]
        ):

            pdf.set_font(
                "Times",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                6,
                line
            )

            pdf.set_font(
                "Times",
                size=11
            )


        elif line.startswith("-"):

            pdf.multi_cell(
                0,
                6,
                "- " + line[1:].strip()
            )


        elif line.startswith("*"):

            pdf.multi_cell(
                0,
                6,
                "- " + line[1:].strip()
            )


        else:

            pdf.multi_cell(
                0,
                6,
                line
            )


    result = pdf.output()

    return bytes(result)
